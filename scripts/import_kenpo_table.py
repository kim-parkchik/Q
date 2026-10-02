#!/usr/bin/env python3
"""
協会けんぽの「健康保険・厚生年金保険の保険料額表（全都道府県分 Excel）」から、
Q で使うデータファイル（TypeScript）を作る。

使い方:
    pip install openpyxl
    python3 scripts/import_kenpo_table.py insurance-tables/協会けんぽ_保険料額表_令和8年度.xlsx 2026 \\
        src/constants/insurance/kenpoTable2026.ts \\
        https://www.kyoukaikenpo.or.jp/g7/cat330/sb3150/r08/r8ryougakuhyou3gatukara/

・数字を手や AI で書き写さないための道具。料率と標準報酬月額表は Excel から機械的に読み取る。
・読み取った料率で計算した保険料が、Excel に印刷されている全額・折半額と全都道府県・全等級で
  一致するかを確かめ、1つでも合わなければエラーで止まる。
"""
import re
import sys
import warnings
from decimal import Decimal, ROUND_HALF_UP

import openpyxl

warnings.filterwarnings("ignore")  # 印刷範囲の警告を出さない


def pct(v):
    """Excel の 0.1027999… を 10.28（%）に直す（小数第3位まで）"""
    return float((Decimal(str(v)) * 100).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP))


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).replace(",", "").strip()
    return float(s) if re.fullmatch(r"[\d.]+", s) else None


def find_rate_row(rows):
    """「等級」「月額」の行（料率が並んでいる行）を探す"""
    for i, r in enumerate(rows):
        if r and str(r[0]).strip() == "等級":
            return i
    raise ValueError("「等級」の行が見つかりません")


def parse_sheet(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    text = "\n".join(" ".join(str(x) for x in r if x not in (None, "")) for r in rows)
    ri = find_rate_row(rows)
    rate_cells = [c for c in rows[ri] if isinstance(c, float) and 0 < c < 1]
    # 並び: 健康(介護なし), 健康(介護あり), [子ども・子育て支援金], 厚生年金
    if len(rate_cells) == 4:
        health, health_care, support, pension = rate_cells
    elif len(rate_cells) == 3:
        health, health_care, pension = rate_cells
        support = None
    else:
        raise ValueError(f"{ws.title}: 料率の数が想定と違います {rate_cells}")

    grades = []
    for r in rows[ri + 1:]:
        if not r or r[0] is None:
            continue
        m = re.fullmatch(r"\s*(\d+)\s*(?:[（(]\s*(\d+)\s*[）)])?\s*", str(r[0]))
        if not m or num(r[1]) is None:
            continue
        g = {
            "grade": int(m.group(1)),
            "pensionGrade": int(m.group(2)) if m.group(2) else None,
            "standard": int(num(r[1])),
            "from": int(num(r[2])) if num(r[2]) is not None else 0,
            "to": int(num(r[4])) if num(r[4]) is not None else None,
            "health_full": num(r[5]), "health_half": num(r[6]),
            "care_full": num(r[7]), "care_half": num(r[8]),
        }
        rest = [num(x) for x in r[9:13]]
        if support is not None:
            g["support_full"], g["support_half"] = rest[0], rest[1]
            g["pension_full"], g["pension_half"] = rest[2], rest[3]
        else:
            g["pension_full"], g["pension_half"] = rest[0], rest[1]
        grades.append(g)
    return {"health": health, "health_care": health_care, "support": support, "pension": pension,
            "grades": grades, "text": text}


def close(a, b):
    return a is not None and abs(a - b) < 0.011


def main():
    src, year, out = sys.argv[1], sys.argv[2], sys.argv[3]
    source_url = sys.argv[4] if len(sys.argv) > 4 else "https://www.kyoukaikenpo.or.jp/g7/cat330/sb3150/"
    wb = openpyxl.load_workbook(src, data_only=True)
    sheets = {name: parse_sheet(wb[name]) for name in wb.sheetnames}
    assert len(sheets) == 47, f"都道府県のシートが47ではありません: {len(sheets)}"
    first = next(iter(sheets.values()))
    title = wb[wb.sheetnames[0]].cell(1, 1).value.strip()
    text = first["text"]

    # 全国共通の値
    care = round(pct(first["health_care"]) - pct(first["health"]), 3)
    pension = pct(first["pension"])
    support = pct(first["support"]) if first["support"] else None
    m = re.search(r"拠出金率（([\d.]+)％）", text)
    assert m, "子ども・子育て拠出金率が見つかりません"
    contribution = float(m.group(1))
    m = re.search(r"([０-９\d]+)年([０-９\d]+)月分（", title)
    eff_month = int(m.group(2).translate(str.maketrans("０１２３４５６７８９", "0123456789")))
    sup_from = None
    if support:
        m = re.search(r"子ども・子育て支援金率：令和(\d+)年(\d+)月分", text)
        assert m, "子ども・子育て支援金の適用開始月が見つかりません"
        sup_from = (2018 + int(m.group(1)), int(m.group(2)))

    grades = [{k: g[k] for k in ("grade", "pensionGrade", "standard", "from", "to")} for g in first["grades"]]
    assert grades[0]["grade"] == 1 and len(grades) >= 50, f"等級の数がおかしい: {len(grades)}"
    for a, b in zip(grades, grades[1:]):
        assert a["to"] == b["from"], f"等級がつながっていません: {a} → {b}"

    # 全都道府県・全等級で、料率から計算した保険料が Excel の印刷値と一致するか確かめる
    health_rates = {}
    checked = 0
    for name, s in sheets.items():
        h = pct(s["health"])
        assert round(pct(s["health_care"]) - h, 3) == care, f"{name}: 介護保険料率が他の県と違います"
        assert pct(s["pension"]) == pension, f"{name}: 厚生年金保険料率が違います"
        assert [{k: g[k] for k in ("grade", "pensionGrade", "standard", "from", "to")} for g in s["grades"]] == grades, f"{name}: 等級表が違います"
        health_rates[name] = h
        for g in s["grades"]:
            std = g["standard"]
            assert close(g["health_full"], std * h / 100), f"{name} {g['grade']}級: 健康保険料が合いません"
            assert close(g["health_half"], std * h / 200), f"{name} {g['grade']}級: 健康保険料（折半）が合いません"
            assert close(g["care_full"], std * (h + care) / 100), f"{name} {g['grade']}級: 介護ありの保険料が合いません"
            if support:
                assert close(g["support_full"], std * support / 100), f"{name} {g['grade']}級: 支援金が合いません"
            if g["pensionGrade"]:
                assert close(g["pension_full"], std * pension / 100), f"{name} {g['grade']}級: 厚生年金保険料が合いません"
            checked += 1

    pension_grades = [g for g in grades if g["pensionGrade"]]
    L = []
    L.append("/**")
    L.append(f" * {title}")
    L.append(" *")
    L.append(" * ⚠️ このファイルは自動生成です。手で書き換えないでください。")
    L.append(" * 元データ: 協会けんぽの Excel（insurance-tables/ フォルダ）")
    L.append(f" * 出典    : 全国健康保険協会ホームページ（{source_url}）の保険料額表を加工して作成")
    L.append(" *           （Excel の表を Q で使える形に変換したもので、協会けんぽが作成したものではありません）")
    L.append(" * 作り方  : python3 scripts/import_kenpo_table.py <Excel> <年度> <このファイル> <出典URL>")
    L.append(f" * 確認    : 全47都道府県 × 全{len(grades)}等級（{checked}行）の保険料が Excel の印刷値と一致することを確認済み")
    L.append(" */")
    L.append("import type { KenpoTable } from './kenpoTableTypes';")
    L.append("")
    L.append(f"export const KENPO_TABLE_{year}: KenpoTable = {{")
    L.append(f'  title: "{title}",')
    L.append(f"  /** この年度の料率を使い始める月（保険料の対象月） */")
    L.append(f"  effectiveFrom: {{ year: {year}, month: {eff_month} }},")
    L.append(f"  /** 介護保険料率（%・全国一律） */")
    L.append(f"  careRate: {care},")
    L.append(f"  /** 厚生年金保険料率（%） */")
    L.append(f"  pensionRate: {pension},")
    L.append(f"  /** 子ども・子育て支援金率（%・全国一律）。制度がない年度は null */")
    L.append(f"  childSupportRate: {support if support else 'null'},")
    L.append(f"  /** 子ども・子育て支援金を徴収し始める月 */")
    L.append(f"  childSupportFrom: {('{ year: %d, month: %d }' % sup_from) if sup_from else 'null'},")
    L.append(f"  /** 子ども・子育て拠出金率（%・事業主のみ負担） */")
    L.append(f"  childCareContributionRate: {contribution},")
    L.append(f"  /** 都道府県別の健康保険料率（%・介護保険第2号被保険者に該当しない場合） */")
    L.append("  healthRates: {")
    for name, h in health_rates.items():
        L.append(f'    "{name}": {h},')
    L.append("  },")
    L.append("  /** 標準報酬月額表（健康保険の等級。pensionGrade は厚生年金の等級、範囲外は null） */")
    L.append("  grades: [")
    for g in grades:
        to = g["to"] if g["to"] is not None else "Infinity"
        pg = g["pensionGrade"] if g["pensionGrade"] else "null"
        L.append(f"    {{ grade: {g['grade']}, pensionGrade: {pg}, standard: {g['standard']}, from: {g['from']}, to: {to} }},")
    L.append("  ],")
    L.append(f"  /** 厚生年金の標準報酬月額の下限・上限 */")
    L.append(f"  pensionMinStandard: {pension_grades[0]['standard']},")
    L.append(f"  pensionMaxStandard: {pension_grades[-1]['standard']},")
    L.append("};")
    open(out, "w", encoding="utf-8").write("\n".join(L) + "\n")
    print(f"{title}: {len(health_rates)}都道府県・{len(grades)}等級、{checked}行の保険料を照合 → {out}")


if __name__ == "__main__":
    main()
